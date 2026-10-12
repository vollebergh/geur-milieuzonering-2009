"""Build the odour-distance vs. environmental-category analysis workbook.

Reads the three source workbooks in ../doc (read-only, never written), extracts
the relevant columns, adds verification formulas, builds value-only datasets,
computes rank correlations and writes:

    analysis/geur_categorie_analyse.xlsx   intermediate + final data, statistics
    analysis/figures/*.png                 correlation diagrams (count heatmaps)

Usage (from the repository root):
    python3 analysis/build_analysis.py
Afterwards recalculate the formulas once with LibreOffice so cached values exist:
    python3 <xlsx-skill>/scripts/recalc.py analysis/geur_categorie_analyse.xlsx
"""

import hashlib
import shutil
import subprocess
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
import xlrd
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "doc"
OUT_DIR = ROOT / "analysis"
FIG_DIR = OUT_DIR / "figures"
OUT_XLSX = OUT_DIR / "geur_categorie_analyse.xlsx"

# Source definitions. Column letters refer to the source sheets.
#   cols: id1/id2/nummer/descr = identifiers; geur/stof/geluid/gevaar = distances;
#         max = 'GROOTSTE AFSTAND'; cat = category column used in the analysis.
SOURCES = {
    "T1": {
        "path": DOC / "bijlage_1_richtafstandenlijst_tabel_1_versie_maart_2009_def.xls",
        "sheet": "EINDTABEL 1 met formule",
        "title": "Bijlage 1, tabel 1 (bedrijfsactiviteiten)",
        "cols": {"id1": "A", "id2": "B", "nummer": "C", "descr": "D", "geur": "E", "stof": "F",
                 "geluid": "G", "gevaar": "J", "max": "O", "cat": "R"},
        "id_labels": ("SBI-1993", "SBI-2008"),
    },
    "T2": {
        "path": DOC / "bijlage_1_richtafstandenlijst_tabel_2_versie_maart_2009_def.xls",
        "sheet": "TABEL2",
        "title": "Bijlage 1, tabel 2 (opslagen en installaties)",
        "cols": {"id1": "A", "id2": None, "nummer": "B", "descr": "C", "geur": "D", "stof": "E",
                 "geluid": "F", "gevaar": "I", "max": "N", "cat": "P"},
        "id_labels": ("Nr.", None),
    },
    "B2": {
        "path": DOC / "bijlage_2_staat_bedrijfsactiviteiten_bedrijventerreinen_versie_maart_2009_def.xlsx",
        "sheet": "EINDTABEL 1 met formule",
        "title": "Bijlage 2 (bedrijventerreinen)",
        "cols": {"id1": "A", "id2": "B", "nummer": "C", "descr": "D", "geur": "E", "stof": "F",
                 "geluid": "G", "gevaar": "J", "max": "O", "cat": "R"},
        "id_labels": ("SBI-CODE", "SBI-CODE 2008"),
    },
}
# Datasets analysed: the three sources plus bijlage 1 as a whole (tabel 1 + tabel 2).
DATASET_TITLES = {
    "T1": SOURCES["T1"]["title"],
    "T2": SOURCES["T2"]["title"],
    "B1": "Bijlage 1, tabel 1 + tabel 2",
    "B2": SOURCES["B2"]["title"],
}
FIGURE_KEYS = ("T1", "T2", "B1", "B2")

# Distance (m) -> category label, as in the 'codetabel' sheet of both sources.
CODETABLE = [
    (10, "1"), (30, "2"), (50, "3.1"), (100, "3.2"), (200, "4.1"),
    (300, "4.2"), (500, "5.1"), (700, "5.2"), (1000, "5.3"), (1500, "6"),
]
CAT_RANK = {label: i + 1 for i, (_, label) in enumerate(CODETABLE)}
DIST_TO_CAT = dict(CODETABLE)
ODOUR_STEPS = [0, 10, 30, 50, 100, 200, 300, 500, 700, 1000, 1500]

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
CHECK_FILL = PatternFill("solid", fgColor="FFF2CC")
FORMULA_FONT = Font(name=FONT, color="000000")
INPUT_FONT = Font(name=FONT, color="0000FF")


# --------------------------------------------------------------------------- helpers

def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def to_num(v):
    """Numeric value of a raw distance cell; None for empty, '-' or text."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def to_int_or_none(v):
    n = to_num(v)
    return None if n is None else int(n) if float(n).is_integer() else n


def cat_label(v):
    """Normalise a cached category value (int 2, float 3.2, str '3.1') to text."""
    if v is None or v == "":
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s if s in CAT_RANK else s


def raw_content(cell_formula_value):
    v = cell_formula_value
    if v is None:
        return None
    return str(v)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def midrank(values):
    """Average ranks (1-based), identical to the COUNTIF midrank formula."""
    return list(stats.rankdata(values, method="average"))


# --------------------------------------------------------------------------- extraction

def xls_to_xlsx(path, tmpdir):
    """Convert a legacy .xls to .xlsx with LibreOffice (only to read formula text)."""
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmpdir), str(path)],
                   check=True, capture_output=True, timeout=180)
    return Path(tmpdir) / (path.stem + ".xlsx")


def load_cells(src, tmpdir):
    """Return (content(col, row), cached(col, row), codetable rows, last row) for a source.

    content = formula text or typed value as stored in the source; cached = the value
    Excel displays. For .xls files the cached values are read with xlrd (Excel's own
    stored results); the formula text comes from a LibreOffice conversion in tmpdir.
    """
    path = src["path"]
    if path.suffix == ".xls":
        book = xlrd.open_workbook(str(path))
        sh = book.sheet_by_name(src["sheet"])
        ct = book.sheet_by_name("codetabel")
        conv = openpyxl.load_workbook(xls_to_xlsx(path, tmpdir))[src["sheet"]]

        def cached(col, r):
            c, rr = openpyxl.utils.column_index_from_string(col) - 1, r - 1
            if rr >= sh.nrows or c >= sh.row_len(rr):
                return None
            if sh.cell_type(rr, c) in (xlrd.XL_CELL_ERROR, xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                return None
            return sh.cell_value(rr, c)

        def content(col, r):
            v = conv[f"{col}{r}"].value
            if isinstance(v, str) and v.startswith("="):
                return v.replace("FALSE()", "FALSE").replace("TRUE()", "TRUE")
            return cached(col, r)

        codetable = [(ct.cell_value(r, 0), ct.cell_value(r, 1)) for r in range(2, 12)]
        last = sh.nrows
    else:
        wf = openpyxl.load_workbook(path)[src["sheet"]]
        wv = openpyxl.load_workbook(path, data_only=True)
        vs = wv[src["sheet"]]
        content = lambda col, r: wf[f"{col}{r}"].value
        cached = lambda col, r: vs[f"{col}{r}"].value
        codetable = [(wv["codetabel"][f"A{r}"].value, wv["codetabel"][f"B{r}"].value) for r in range(3, 13)]
        last = wf.max_row
    return content, cached, codetable, last


def read_source(key, src, tmpdir):
    content, cached, codetable, last = load_cells(src, tmpdir)
    found = [(int(d), cat_label(c)) for d, c in codetable]
    assert found == CODETABLE, f"{key}: codetabel differs: {found}"

    c = src["cols"]
    keycols = [c[k] for k in ("id1", "id2", "nummer", "descr", "geur", "stof", "geluid", "gevaar",
                              "max", "cat") if c[k]]
    last = max(r for r in range(1, last + 1) if any(content(col, r) not in (None, "") for col in keycols))
    rows = []
    for r in range(3, last + 1):
        if all(content(col, r) in (None, "") for col in keycols):
            continue
        g = lambda k: content(c[k], r) if c[k] else None
        descr = g("descr")
        descr = (" ".join(descr.split()) if isinstance(descr, str) else descr) or None
        num = lambda v: int(v) if isinstance(v, float) and v.is_integer() else v
        cat_content = g("cat")
        rows.append({
            "src_row": r,
            "sbi_old": num(g("id1")), "sbi_2008": g("id2"), "nummer": num(g("nummer")), "descr": descr,
            "odour_raw": num(g("geur")), "dust_raw": num(g("stof")),
            "noise_raw": num(g("geluid")), "danger_raw": num(g("gevaar")),
            "O_content": raw_content(num(g("max"))), "O_cached": to_int_or_none(cached(c["max"], r)),
            "R_content": raw_content(num(cat_content)), "R_cached": cat_label(cached(c["cat"], r)),
            "R_is_formula": isinstance(cat_content, str) and cat_content.startswith("="),
        })
    return rows


def derive(row):
    """Python replica of the verification formulas in the extract sheets."""
    nums = [to_int_or_none(row[k]) for k in ("odour_raw", "dust_raw", "noise_raw", "danger_raw")]
    odour, others = nums[0], [n for n in nums[1:] if n is not None]
    present = [n for n in nums if n is not None]
    max_rec = max(present) if present else None
    cat_rec = DIST_TO_CAT.get(max_rec) if max_rec is not None else None
    reason = None
    if not row["descr"]:
        reason = "no activity description"
    elif odour is None:
        reason = "no numeric odour distance"
    elif not row["R_cached"]:
        reason = "no category"
    other_max = max(others) if others else 0
    return {
        "odour": odour, "max_rec": max_rec, "cat_rec": cat_rec,
        "O_ok": (row["O_cached"] == max_rec),
        "R_ok": (row["R_cached"] or None) == cat_rec,
        "included": reason is None, "reason": reason,
        "decisive": (odour == max_rec) if reason is None else None,
        "other_max": other_max,
        "cat_rank_excl": CAT_RANK.get(DIST_TO_CAT.get(other_max), 0),
    }


# --------------------------------------------------------------------------- workbook sheets

EXTRACT_HEADERS = [
    "Source row", "{id1}", "{id2}", "Nummer", "Omschrijving",
    "Geur ({geur}) raw", "Stof ({stof}) raw", "Geluid ({geluid}) raw", "Gevaar ({gevaar}) raw",
    "Max ({max}) source content", "Max ({max}) cached value", "Category ({cat}) source content",
    "Category ({cat}) cached value", "Category is formula",
    # check block (formulas)
    "Geur num", "Stof num", "Geluid num", "Gevaar num",
    "Max distance recalc", "Category recalc", "O check", "R check",
    "Exclusion reason", "Included", "Category rank", "Geur decisive",
]


def write_extract(wb, key, rows):
    ws = wb.create_sheet(f"{key}_extract")
    src = SOURCES[key]
    fmt = dict(src["cols"], id1=src["id_labels"][0], id2=src["id_labels"][1] or "(none)")
    ws.append([h.format(**fmt) for h in EXTRACT_HEADERS])
    style_header(ws, 1, len(EXTRACT_HEADERS))
    for i, row in enumerate(rows, start=2):
        ws.append([
            row["src_row"], row["sbi_old"], row["sbi_2008"], row["nummer"], row["descr"],
            row["odour_raw"], row["dust_raw"], row["noise_raw"], row["danger_raw"],
            row["O_content"], row["O_cached"], row["R_content"], row["R_cached"], row["R_is_formula"],
        ])
        n = i
        for col, src in zip("OPQR", "FGHI"):
            ws[f"{col}{n}"] = f'=IF({src}{n}="","",IF(ISNUMBER({src}{n}),{src}{n},IFERROR(VALUE({src}{n}),"")))'
        ws[f"S{n}"] = f'=IF(COUNT(O{n}:R{n})=0,"",MAX(O{n}:R{n}))'
        ws[f"T{n}"] = f'=IF(S{n}="","",IFERROR(INDEX(Codetabel!$B$2:$B$11,MATCH(S{n},Codetabel!$A$2:$A$11,0)),""))'
        ws[f"U{n}"] = f'=IF(NOT(X{n}),"n/a (excluded)",IF(K{n}=S{n},"OK","MISMATCH"))'
        ws[f"V{n}"] = f'=IF(NOT(X{n}),"n/a (excluded)",IF(M{n}&""=T{n}&"","OK","MISMATCH"))'
        ws[f"W{n}"] = (f'=IF(E{n}="","no activity description",IF(O{n}="","no numeric odour distance",'
                       f'IF(M{n}="","no category","")))')
        ws[f"X{n}"] = f'=W{n}=""'
        ws[f"Y{n}"] = f'=IF(X{n},INDEX(Codetabel!$C$2:$C$11,MATCH(M{n}&"",Codetabel!$B$2:$B$11,0)),"")'
        ws[f"Z{n}"] = f'=IF(X{n},O{n}=S{n},"")'
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = INPUT_FONT if cell.column <= 14 else FORMULA_FONT
            if cell.column > 14:
                cell.fill = CHECK_FILL
    set_widths(ws, [8, 10, 12, 8, 50, 9, 9, 9, 9, 26, 9, 22, 9, 9,
                    8, 8, 8, 8, 10, 10, 10, 10, 24, 9, 9, 9])
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:Z{len(rows) + 1}"


DATA_HEADERS = [
    "Source", "Source row", "SBI (old)", "SBI-2008", "Omschrijving",
    "Geur (m)", "Categorie", "Category rank", "Max distance (m)", "Geur decisive",
    "Max other distances (m)", "Category rank excl. geur", "Also in bijlage 2 / tabel 1",
    # formulas
    "Midrank geur", "Midrank category", "Midrank geur (not decisive)",
    "Midrank category (not decisive)", "Midrank category excl. geur",
    "Geur step index (0-10)", "Category step index (0-9)",
]


def write_dataset(wb, name, records):
    ws = wb.create_sheet(name)
    ws.append(DATA_HEADERS)
    style_header(ws, 1, len(DATA_HEADERS))
    last = len(records) + 1
    F, H, J, L = (f"${c}$2:${c}${last}" for c in "FHJL")
    for i, rec in enumerate(records, start=2):
        ws.append([rec[k] for k in ("source", "src_row", "sbi_old", "sbi_2008", "descr", "odour",
                                    "cat", "cat_rank", "max_rec", "decisive", "other_max",
                                    "cat_rank_excl", "in_both")])
        ws[f"N{i}"] = f'=COUNTIF({F},"<"&F{i})+(COUNTIF({F},F{i})+1)/2'
        ws[f"O{i}"] = f'=COUNTIF({H},"<"&H{i})+(COUNTIF({H},H{i})+1)/2'
        ws[f"P{i}"] = (f'=IF(J{i}=FALSE,COUNTIFS({J},FALSE,{F},"<"&F{i})'
                       f'+(COUNTIFS({J},FALSE,{F},F{i})+1)/2,"")')
        ws[f"Q{i}"] = (f'=IF(J{i}=FALSE,COUNTIFS({J},FALSE,{H},"<"&H{i})'
                       f'+(COUNTIFS({J},FALSE,{H},H{i})+1)/2,"")')
        ws[f"R{i}"] = f'=COUNTIF({L},"<"&L{i})+(COUNTIF({L},L{i})+1)/2'
        ws[f"S{i}"] = f'=MATCH(F{i},Codetabel!$G$2:$G$12,0)-1'
        ws[f"T{i}"] = f'=H{i}-1'
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = INPUT_FONT if cell.column <= 13 else FORMULA_FONT
    set_widths(ws, [7, 8, 10, 12, 50, 8, 9, 9, 10, 9, 11, 11, 9, 9, 9, 11, 11, 11, 10, 10])
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:T{last}"
    return last


def write_codetable(wb):
    ws = wb.create_sheet("Codetabel")
    ws.append(["Distance (m)", "Category", "Ordinal rank"])
    style_header(ws, 1, 3)
    for d, lab in CODETABLE:
        ws.append([d, lab, CAT_RANK[lab]])
    ws["E1"] = "Source: sheet 'codetabel' in all three source workbooks (identical; verified by build_analysis.py)."
    ws["E2"] = "Category labels stored as text; rank 1-10 is the ordinal position used in the rank correlations."
    ws["E3"] = "Rank 0 (only in 'Category rank excl. geur') means all non-odour distances are 0 m."
    ws["G1"] = "Geur steps (m)"
    ws["H1"] = "Step index"
    for k, step in enumerate(ODOUR_STEPS):
        ws.cell(row=2 + k, column=7, value=step)
        ws.cell(row=2 + k, column=8, value=k)
    ws["E5"] = "G:H = the 11 distinct geur distances and their step index, used for the regression lines (sheet 'Regression')."
    set_widths(ws, [13, 10, 12, 2, 70, 2, 14, 10])
    style_header(ws, 1, 3)
    ws["G1"].font = ws["H1"].font = Font(name=FONT, bold=True)


def write_crosstab(wb, name, data_sheet, last, title):
    ws = wb.create_sheet(name)
    ws["A1"] = f"Cross-tabulation geur distance x category - {title}"
    ws["A1"].font = Font(name=FONT, bold=True)
    ws["A2"] = ("Counts via COUNTIFS on sheet " + data_sheet + ". Cells right of the bold diagonal are "
                "structurally impossible: category = max(geur, stof, geluid, gevaar), so geur <= category distance.")
    hdr = ["Category \\ geur (m)"] + ODOUR_STEPS + ["Total", "Mean geur (m)"]
    for c, h in enumerate(hdr, start=1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, 4, len(hdr))
    rng = lambda col: f"'{data_sheet}'!${col}$2:${col}${last}"
    for i, (dist, lab) in enumerate(CODETABLE):
        r = 5 + i
        ws.cell(row=r, column=1, value=lab)
        for j, step in enumerate(ODOUR_STEPS):
            cell = ws.cell(row=r, column=2 + j,
                           value=f'=COUNTIFS({rng("G")},$A{r},{rng("F")},{get_column_letter(2 + j)}$4)')
            if step == dist:
                cell.font = Font(name=FONT, bold=True)
        tcol = 2 + len(ODOUR_STEPS)
        ws.cell(row=r, column=tcol, value=f"=SUM(B{r}:{get_column_letter(tcol - 1)}{r})")
        ws.cell(row=r, column=tcol + 1,
                value=f'=IFERROR(AVERAGEIFS({rng("F")},{rng("G")},$A{r}),"")')
    tr = 5 + len(CODETABLE)
    ws.cell(row=tr, column=1, value="Total")
    for j in range(len(ODOUR_STEPS) + 1):
        col = get_column_letter(2 + j)
        ws.cell(row=tr, column=2 + j, value=f"=SUM({col}5:{col}{tr - 1})")
    ws.cell(row=tr + 1, column=1, value="Check: total = rows in dataset")
    ws.cell(row=tr + 1, column=2,
            value=f'=IF({get_column_letter(2 + len(ODOUR_STEPS))}{tr}=COUNT({rng("F")}),"OK","MISMATCH")')
    ws.cell(row=tr + 1, column=2).fill = CHECK_FILL
    set_widths(ws, [20] + [7] * len(ODOUR_STEPS) + [8, 13])


# --------------------------------------------------------------------------- statistics

def corr_stats(x, y):
    rho, p_rho = stats.spearmanr(x, y)
    tau, p_tau = stats.kendalltau(x, y, variant="b")
    r, p_r = stats.pearsonr(x, y)
    return {"n": len(x), "rho": rho, "p_rho": p_rho, "tau": tau, "p_tau": p_tau, "r": r, "p_r": p_r}


def bootstrap_ci(x, y, n_boot=2000, seed=1):
    import random
    rnd = random.Random(seed)
    idx = list(range(len(x)))
    vals = []
    for _ in range(n_boot):
        s = [rnd.choice(idx) for _ in idx]
        xs, ys = [x[k] for k in s], [y[k] for k in s]
        if len(set(xs)) > 1 and len(set(ys)) > 1:
            vals.append(stats.spearmanr(xs, ys)[0])
    vals.sort()
    return vals[int(0.025 * len(vals))], vals[int(0.975 * len(vals)) - 1]


def permutation_baseline(recs, n_perm=1000, seed=1):
    """Mean Spearman rho if geur were unrelated to the other distances.

    Geur values are shuffled across activities; the category is then rebuilt as
    max(shuffled geur, max other distances) via the code table. The mean rho over
    the permutations is the correlation produced by the max() construction alone.
    """
    import random
    rnd = random.Random(seed)
    odour = [r["odour"] for r in recs]
    other = [r["other_max"] for r in recs]
    rhos = []
    for _ in range(n_perm):
        perm = odour[:]
        rnd.shuffle(perm)
        cat = [CAT_RANK.get(DIST_TO_CAT.get(max(g, o)), 0) for g, o in zip(perm, other)]
        rhos.append(stats.spearmanr(perm, cat)[0])
    rhos.sort()
    return sum(rhos) / len(rhos), rhos[int(0.025 * n_perm)], rhos[int(0.975 * n_perm) - 1]


def write_statistics(wb, analyses, extract_checks, src_hashes):
    ws = wb.create_sheet("Statistics", 1)
    ws["A1"] = "Rank correlation between geur distance and environmental category"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    ws["A2"] = ("Black = live formula on the dataset sheets (re-computes in Excel/LibreOffice). "
                "Blue = value computed by build_analysis.py (scipy 'spearmanr', 'kendalltau' variant b, "
                "'pearsonr'; 2000-sample percentile bootstrap, seed 1).")
    hdr = ["#", "Analysis", "Dataset sheet", "n (formula)", "Spearman rho (formula)",
           "Spearman rho (Python)", "Check", "Spearman 95% CI (bootstrap)", "p (Spearman)",
           "Kendall tau-b (Python)", "p (Kendall)", "Pearson r geur m vs rank (formula)",
           "Pearson r (Python)", "p (Pearson)",
           "Spearman rho from max() construction alone (permutation mean, 95% range)"]
    for c, h in enumerate(hdr, start=1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, 4, len(hdr))
    for i, a in enumerate(analyses, start=1):
        r = 4 + i
        s, last = a["sheet"], a["last"]
        xc, yc = a["rank_cols"]
        xr, yr = f"'{s}'!{xc}2:{xc}{last}", f"'{s}'!{yc}2:{yc}{last}"
        rawx, rawy = a["raw_cols"]
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=a["label"])
        ws.cell(row=r, column=3, value=s)
        ws.cell(row=r, column=4, value=f"=COUNT({xr})")
        ws.cell(row=r, column=5, value=f"=CORREL({xr},{yr})")
        st = a["stats"]
        ws.cell(row=r, column=6, value=round(st["rho"], 6))
        ws.cell(row=r, column=7, value=f'=IF(ABS(E{r}-F{r})<0.0001,"OK","MISMATCH")')
        ws.cell(row=r, column=8, value=f"{a['ci'][0]:.3f} to {a['ci'][1]:.3f}")
        ws.cell(row=r, column=9, value=float(f"{st['p_rho']:.3g}"))
        ws.cell(row=r, column=10, value=round(st["tau"], 6))
        ws.cell(row=r, column=11, value=float(f"{st['p_tau']:.3g}"))
        if a.get("subset_filter") is None:
            ws.cell(row=r, column=12, value=f"=CORREL('{s}'!{rawx}2:{rawx}{last},'{s}'!{rawy}2:{rawy}{last})")
        else:
            ws.cell(row=r, column=12, value="n/a (subset)")
        ws.cell(row=r, column=13, value=round(st["r"], 6))
        ws.cell(row=r, column=14, value=float(f"{st['p_r']:.3g}"))
        pb = a.get("perm")
        ws.cell(row=r, column=15,
                value=f"{pb[0]:.3f} ({pb[1]:.3f} to {pb[2]:.3f})" if pb else "n/a")
        for c in range(1, len(hdr) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = INPUT_FONT if c in (6, 8, 9, 10, 11, 13, 14, 15) else FORMULA_FONT
            if c in (5, 6, 10, 12, 13):
                cell.number_format = "0.000"
            if c in (9, 11, 14):
                cell.number_format = "0.0E+00"
        ws.cell(row=r, column=7).fill = CHECK_FILL

    r0 = 6 + len(analyses)
    ws.cell(row=r0, column=1, value="Verification of source values").font = Font(name=FONT, bold=True)
    hdr2 = ["Source", "Rows extracted", "Rows included", "O cached = recalculated (formula)",
            "R cached = recalculated (formula)", "R hardcoded (not formula) among included",
            "Geur decisive among included", "Source file: SHA-256 (unchanged)"]
    for c, h in enumerate(hdr2, start=1):
        ws.cell(row=r0 + 1, column=c, value=h)
    style_header(ws, r0 + 1, len(hdr2))
    for k, (key, chk) in enumerate(extract_checks.items()):
        r = r0 + 2 + k
        sh, n = f"{key}_extract", chk["n"] + 1
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=f"=COUNTA('{sh}'!A2:A{n})")
        ws.cell(row=r, column=3, value=f"=COUNTIF('{sh}'!X2:X{n},TRUE)")
        ws.cell(row=r, column=4,
                value=f"=COUNTIFS('{sh}'!X2:X{n},TRUE,'{sh}'!U2:U{n},\"OK\")&\" of \"&C{r}")
        ws.cell(row=r, column=5,
                value=f"=COUNTIFS('{sh}'!X2:X{n},TRUE,'{sh}'!V2:V{n},\"OK\")&\" of \"&C{r}")
        ws.cell(row=r, column=6, value=f"=COUNTIFS('{sh}'!X2:X{n},TRUE,'{sh}'!N2:N{n},FALSE)")
        ws.cell(row=r, column=7, value=f"=COUNTIF('{sh}'!Z2:Z{n},TRUE)")
        ws.cell(row=r, column=8, value=f"{SOURCES[key]['path'].name}: {src_hashes[key]}")
        for c in range(1, 9):
            ws.cell(row=r, column=c).font = INPUT_FONT if c == 8 else FORMULA_FONT
    set_widths(ws, [4, 58, 12, 10, 12, 12, 10, 18, 11, 12, 11, 14, 12, 11, 24])
    return ws


# --------------------------------------------------------------------------- figure

def heatmap(ax, records, title, swap=False):
    """Count heatmap; swap=False: x = geur, y = category; swap=True: x = category, y = geur."""
    cats = [lab for _, lab in CODETABLE]
    grid = [[0] * len(ODOUR_STEPS) for _ in cats]  # grid[category][geur]
    for rec in records:
        grid[cats.index(rec["cat"])][ODOUR_STEPS.index(rec["odour"])] += 1
    if swap:
        grid = [list(col) for col in zip(*grid)]  # grid[geur][category]
    # cell(i, j) -> (category index, geur index) for row i, column j
    cell = (lambda i, j: (j, i)) if swap else (lambda i, j: (i, j))
    vmax = max(max(row) for row in grid)
    cmap = matplotlib.colors.LinearSegmentedColormap.from_list(
        "seq", ["#fcfcfb", "#cde2fb", "#86b6ef", "#3987e5", "#1c5cab", "#0d366b"])
    ax.imshow(grid, cmap=cmap, vmin=0, vmax=vmax, origin="lower", aspect="auto")
    for i, row in enumerate(grid):
        for j, v in enumerate(row):
            ci, gi = cell(i, j)
            if v:
                ax.text(j, i, str(v), ha="center", va="center", fontsize=8,
                        color="#ffffff" if v > 0.45 * vmax else "#0b0b0b")
            if ODOUR_STEPS[gi] > CODETABLE[ci][0]:
                ax.add_patch(plt.Rectangle((j - 0.5, i - 0.5), 1, 1, fill=False, hatch="///",
                                           edgecolor="#d6d5d0", linewidth=0))
    # Diagonal: odour distance equals category distance (odour is decisive).
    for ci, (dist, _) in enumerate(CODETABLE):
        gi = ODOUR_STEPS.index(dist)
        x, y = (ci, gi) if swap else (gi, ci)
        ax.add_patch(plt.Rectangle((x - 0.5, y - 0.5), 1, 1, fill=False,
                                   edgecolor="#52514e", linewidth=1.2))
    odour_ticks = [str(s) for s in ODOUR_STEPS]
    xt, yt = (cats, odour_ticks) if swap else (odour_ticks, cats)
    xl, yl = ("Milieucategorie", "Geurafstand (m)") if swap else ("Geurafstand (m)", "Milieucategorie")
    ax.set_xticks(range(len(xt)), xt, fontsize=8)
    ax.set_yticks(range(len(yt)), yt, fontsize=8)
    ax.set_xlabel(xl, fontsize=9, color="#52514e")
    ax.set_ylabel(yl, fontsize=9, color="#52514e")
    ax.set_title(title, fontsize=10, loc="left", color="#0b0b0b")
    for sp in ax.spines.values():
        sp.set_visible(False)
    ax.tick_params(length=0, colors="#52514e")


def make_figure(datasets, analyses_by_key, path, swap=False):
    fig, axes = plt.subplots(2, 2, figsize=(13, 11), dpi=150)
    fig.patch.set_facecolor("#fcfcfb")
    for ax, key in zip(axes.flat, FIGURE_KEYS):
        st = analyses_by_key[key]
        heatmap(ax, datasets[key],
                f"{DATASET_TITLES[key]}\nn = {st['n']}, Spearman rho = {st['rho']:.2f}, "
                f"Kendall tau-b = {st['tau']:.2f}", swap=swap)
    fig.text(0.01, 0.01,
             "Cell value = number of activities. Outlined cells: geur distance equals the category "
             "distance (geur decisive). Hatched cells are impossible by construction "
             "(category = max of geur, stof, geluid, gevaar).",
             fontsize=8, color="#52514e")
    fig.tight_layout(rect=(0, 0.025, 1, 1))
    fig.savefig(path, facecolor=fig.get_facecolor())
    plt.close(fig)


def step_regression(records, swap):
    """OLS on step indices. swap=False: x = geur step, y = category step; True: reversed."""
    g = [ODOUR_STEPS.index(r["odour"]) for r in records]
    c = [r["cat_rank"] - 1 for r in records]
    x, y = (c, g) if swap else (g, c)
    return stats.linregress(x, y)


def bubble(ax, records, title, swap=False):
    cats = [lab for _, lab in CODETABLE]
    counts = {}
    for rec in records:
        k = (ODOUR_STEPS.index(rec["odour"]), cats.index(rec["cat"]))
        counts[k] = counts.get(k, 0) + 1
    # Impossible region (geur > category distance), lightly hatched.
    for gi, step in enumerate(ODOUR_STEPS):
        for ci, (dist, _) in enumerate(CODETABLE):
            if step > dist:
                x, y = (ci, gi) if swap else (gi, ci)
                ax.add_patch(plt.Rectangle((x - 0.5, y - 0.5), 1, 1, fill=False, hatch="///",
                                           edgecolor="#e4e3de", linewidth=0, zorder=0))
    scale = 18  # marker area (pt^2) per activity
    xs, ys, ss, ns = [], [], [], []
    for (gi, ci), n in counts.items():
        x, y = (ci, gi) if swap else (gi, ci)
        xs.append(x); ys.append(y); ss.append(n * scale); ns.append(n)
    ax.scatter(xs, ys, s=ss, color="#3987e5", alpha=0.75, edgecolors="#fcfcfb",
               linewidths=1.5, zorder=2)
    for x, y, n in zip(xs, ys, ns):
        if n >= 10:
            ax.text(x, y, str(n), ha="center", va="center", fontsize=7, color="#ffffff", zorder=3)
    reg = step_regression(records, swap)
    nx = len(cats) if swap else len(ODOUR_STEPS)
    lx = [-0.5, nx - 0.5]
    ax.plot(lx, [reg.intercept + reg.slope * v for v in lx], color="#0b0b0b", linewidth=2,
            zorder=4, label=f"OLS: y = {reg.intercept:.2f} + {reg.slope:.2f}x  (R\u00b2 = {reg.rvalue ** 2:.2f})")
    odour_ticks = [str(v) for v in ODOUR_STEPS]
    xt, yt = (cats, odour_ticks) if swap else (odour_ticks, cats)
    xl, yl = ("Milieucategorie", "Geurafstand (m)") if swap else ("Geurafstand (m)", "Milieucategorie")
    ax.set_xticks(range(len(xt)), xt, fontsize=8)
    ax.set_yticks(range(len(yt)), yt, fontsize=8)
    ax.set_xlim(-0.5, len(xt) - 0.5)
    ax.set_ylim(-0.5, len(yt) - 0.5)
    ax.set_xlabel(xl, fontsize=9, color="#52514e")
    ax.set_ylabel(yl, fontsize=9, color="#52514e")
    ax.set_title(title, fontsize=10, loc="left", color="#0b0b0b")
    ax.grid(color="#ecebe7", linewidth=0.6, zorder=0)
    ax.set_axisbelow(True)
    for sp in ax.spines.values():
        sp.set_visible(False)
    ax.tick_params(length=0, colors="#52514e")
    # Size key + regression line in one legend.
    handles = [ax.scatter([], [], s=n * scale, color="#3987e5", alpha=0.75, edgecolors="#fcfcfb",
                          label=f"{n} activities") for n in (5, 20, 40)]
    line = ax.get_lines()[0]
    ax.legend(handles=handles + [line], loc="upper center", bbox_to_anchor=(0.5, -0.11),
              ncol=4, fontsize=8, frameon=False, columnspacing=1.6, handletextpad=0.6)
    return reg


def make_bubble_figure(datasets, analyses_by_key, path, swap=False):
    fig, axes = plt.subplots(2, 2, figsize=(13, 13), dpi=150)
    fig.patch.set_facecolor("#fcfcfb")
    for ax, key in zip(axes.flat, FIGURE_KEYS):
        st = analyses_by_key[key]
        bubble(ax, datasets[key],
               f"{DATASET_TITLES[key]}\nn = {st['n']}, Spearman rho = {st['rho']:.2f}", swap=swap)
    fig.text(0.01, 0.01,
             "Circle area is proportional to the number of activities (counts >= 10 printed). "
             "Regression line: ordinary least squares on the step positions of both axes\n"
             "(geur steps 0-10, categories 0-9), each activity weighted equally. "
             "Hatched cells are impossible by construction.",
             fontsize=8, color="#52514e")
    fig.tight_layout(rect=(0, 0.035, 1, 1))
    fig.savefig(path, facecolor=fig.get_facecolor())
    plt.close(fig)


def write_regression(wb, datasets, lasts, data_sheets):
    ws = wb.create_sheet("Regression", 3)
    ws["A1"] = "Regression lines in the bubble diagrams (OLS on step positions)"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    ws["A2"] = ("x and y are step positions, not metres: geur 0, 10, 30, ... 1500 m = steps 0-10; "
                "category 1 ... 6 = steps 0-9 (columns S and T on the dataset sheets). "
                "Black = live formula (SLOPE/INTERCEPT/RSQ); blue = scipy.stats.linregress.")
    hdr = ["Dataset", "x", "y", "n", "Slope (formula)", "Intercept (formula)", "R\u00b2 (formula)",
           "Slope (Python)", "Intercept (Python)", "R\u00b2 (Python)", "SE slope (Python)",
           "p slope (Python)", "Check"]
    for c, h in enumerate(hdr, start=1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, 4, len(hdr))
    r = 5
    for key in FIGURE_KEYS:
        sh, last = data_sheets[key], lasts[key]
        g, c = f"'{sh}'!S2:S{last}", f"'{sh}'!T2:T{last}"
        for swap in (False, True):
            reg = step_regression(datasets[key], swap)
            xr, yr = (c, g) if swap else (g, c)
            vals = [key, "category step" if swap else "geur step", "geur step" if swap else "category step",
                    f"=COUNT({xr})", f"=SLOPE({yr},{xr})", f"=INTERCEPT({yr},{xr})", f"=RSQ({yr},{xr})",
                    round(reg.slope, 6), round(reg.intercept, 6), round(reg.rvalue ** 2, 6),
                    round(reg.stderr, 6), float(f"{reg.pvalue:.3g}"),
                    f'=IF(AND(ABS(E{r}-H{r})<0.0001,ABS(F{r}-I{r})<0.0001),"OK","MISMATCH")']
            for ci, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=ci, value=v)
                cell.font = INPUT_FONT if 8 <= ci <= 12 else FORMULA_FONT
                if 5 <= ci <= 11:
                    cell.number_format = "0.000"
                if ci == 12:
                    cell.number_format = "0.0E+00"
            ws.cell(row=r, column=13).fill = CHECK_FILL
            r += 1
    set_widths(ws, [9, 15, 15, 7, 11, 11, 10, 11, 11, 10, 11, 11, 9])


# --------------------------------------------------------------------------- main

def main():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    hashes_before = {k: sha256(src["path"]) for k, src in SOURCES.items()}

    extracts, datasets, checks = {}, {}, {}
    with tempfile.TemporaryDirectory() as tmpdir:
        for key, src in SOURCES.items():
            extracts[key] = read_source(key, src, tmpdir)
    for key, rows in extracts.items():
        recs = []
        for row in rows:
            d = derive(row)
            row["_d"] = d
            if d["included"]:
                recs.append({
                    "source": key, "src_row": row["src_row"], "sbi_old": row["sbi_old"],
                    "sbi_2008": row["sbi_2008"], "descr": row["descr"], "odour": d["odour"],
                    "cat": row["R_cached"], "cat_rank": CAT_RANK[row["R_cached"]],
                    "max_rec": d["max_rec"], "decisive": d["decisive"], "other_max": d["other_max"],
                    "cat_rank_excl": d["cat_rank_excl"],
                    "key": (key if key == "T2" else "SBI", str(row["sbi_old"]).strip(),
                            str(row["nummer"]).strip(), row["descr"]),
                })
        datasets[key] = recs
        checks[key] = {"n": len(rows)}

    # Overlap: bijlage 2 is (almost) a subset of tabel 1.
    keys_t1 = {r["key"] for r in datasets["T1"]}
    keys_b2 = {r["key"] for r in datasets["B2"]}
    for key, recs in datasets.items():
        other = keys_t1 if key == "B2" else keys_b2
        for r in recs:
            r["in_both"] = r["key"] in other
    datasets["B1"] = [dict(r) for r in datasets["T1"]] + [dict(r) for r in datasets["T2"]]
    union = [dict(r) for r in datasets["B1"]] + [dict(r) for r in datasets["B2"] if not r["in_both"]]

    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"

    analyses, by_key = [], {}
    data_sheets = {"T1": "Data_T1", "T2": "Data_T2", "B1": "Data_B1", "B2": "Data_B2",
                   "B1uB2": "Data_B1_union_B2"}
    lasts = {}
    for k in ("T1", "T2", "B1", "B2"):
        lasts[k] = write_dataset(wb, data_sheets[k], datasets[k])
    lasts["B1uB2"] = write_dataset(wb, data_sheets["B1uB2"], union)

    def add(label, k, recs, xkey, ykey, rank_cols, raw_cols, subset=None, perm=False):
        x = [r[xkey] for r in recs]
        y = [r[ykey] for r in recs]
        st = corr_stats(x, y)
        a = {"label": label, "sheet": data_sheets[k], "last": lasts[k], "rank_cols": rank_cols,
             "raw_cols": raw_cols, "stats": st, "ci": bootstrap_ci(x, y), "subset_filter": subset,
             "perm": permutation_baseline(recs) if perm else None}
        analyses.append(a)
        return st

    for k in ("B1", "T1", "T2", "B2"):
        recs, nm = datasets[k], DATASET_TITLES[k]
        by_key[k] = add(f"{nm}: geur vs category (all activities)", k, recs,
                        "odour", "cat_rank", ("N", "O"), ("F", "H"), perm=True)
        add(f"{nm}: geur vs category, only activities where geur is NOT decisive", k,
            [r for r in recs if not r["decisive"]], "odour", "cat_rank", ("P", "Q"), ("F", "H"),
            subset="not decisive")
        add(f"{nm}: geur vs category based on stof/geluid/gevaar only (geur removed)", k, recs,
            "odour", "cat_rank_excl", ("N", "R"), ("F", "L"))
    add("Bijlage 1 (tabel 1 + 2) union bijlage 2 (duplicates removed): geur vs category", "B1uB2",
        union, "odour", "cat_rank", ("N", "O"), ("F", "H"), perm=True)

    hashes_mid = {k: sha256(src["path"]) for k, src in SOURCES.items()}
    write_statistics(wb, analyses, checks, hashes_mid)
    for key in SOURCES:
        write_extract(wb, key, extracts[key])
    write_codetable(wb)
    for key in FIGURE_KEYS:
        write_crosstab(wb, f"Crosstab_{key}", data_sheets[key], lasts[key], DATASET_TITLES[key])

    fig_ws = wb.create_sheet("Figure", 2)
    figures = (("correlatiediagram_geur_categorie.png", make_figure, False),
               ("correlatiediagram_categorie_geur.png", make_figure, True),
               ("bubbel_geur_categorie.png", make_bubble_figure, False),
               ("bubbel_categorie_geur.png", make_bubble_figure, True))
    anchor_row = 1
    for name, maker, swap in figures:
        fp = FIG_DIR / name
        maker(datasets, by_key, fp, swap=swap)
        im = XLImage(str(fp))
        im.width, im.height = im.width * 0.45, im.height * 0.45
        fig_ws.add_image(im, f"A{anchor_row}")
        anchor_row += int(im.height / 20) + 3
    write_regression(wb, datasets, lasts, data_sheets)

    t2_mismatch = [r for r in extracts["T2"] if r["_d"]["included"]
                   and not (r["_d"]["O_ok"] and r["_d"]["R_ok"])]
    lines = [
        ("Odour distance (geur) vs. environmental category - analysis workbook", True),
        ("Generated by analysis/build_analysis.py; method in analysis/METHODOLOGY.md, results in analysis/FINDINGS.md.", False),
        ("Source files in doc/ are read only; their SHA-256 hashes are listed on 'Statistics' and are checked before and after the build.", False),
        ("", False),
        ("Sources", True),
        ("T1 = doc/bijlage_1_richtafstandenlijst_tabel_1_versie_maart_2009_def.xls (VNG 2009, bijlage 1 tabel 1: bedrijfsactiviteiten).", False),
        ("T2 = doc/bijlage_1_richtafstandenlijst_tabel_2_versie_maart_2009_def.xls (VNG 2009, bijlage 1 tabel 2: opslagen en installaties).", False),
        ("B2 = doc/bijlage_2_staat_bedrijfsactiviteiten_bedrijventerreinen_versie_maart_2009_def.xlsx (VNG 2009, bijlage 2: bedrijventerreinen).", False),
        ("B1 = T1 + T2 (bijlage 1 as a whole). B1 union B2 = B1 plus the B2 activities that are not in T1.", False),
        ("", False),
        ("Sheets", True),
        ("Statistics - correlation coefficients (formula and Python, cross-checked) and source verification summary.", False),
        ("Figure - correlation diagrams, also in analysis/figures/: count heatmaps (geur on x, then axes reversed) and bubble diagrams with regression line (geur on x, then axes reversed); panels T1, T2, B1, B2.", False),
        ("Regression - slope, intercept and R2 of the regression lines in the bubble diagrams (formula and Python, cross-checked).", False),
        ("Data_T1, Data_T2, Data_B1, Data_B2, Data_B1_union_B2 - FINAL datasets: values only, one row per activity with a numeric geur distance and a category. Columns N-R are rank formulas used by 'Statistics'; S-T are step indices used by 'Regression'.", False),
        ("T1_extract, T2_extract, B2_extract - INTERMEDIATE: every non-empty source row with raw values, the original formula/content of the max-distance and category columns, the cached (displayed) values, and recalculation checks.", False),
        ("Codetabel - distance-to-category lookup with ordinal ranks.", False),
        ("Crosstab_T1, Crosstab_T2, Crosstab_B1, Crosstab_B2 - frequency tables (COUNTIFS) behind the figures.", False),
        ("", False),
        ("Colour conventions", True),
        ("Blue text = value copied from a source file or computed in Python. Black text = live formula. Yellow fill = verification column.", False),
        ("", False),
        ("Column mapping", True),
        ("T1 and B2: geur = E, stof = F, geluid = G, gevaar = J, grootste afstand = O, categorie = R ('CATEGORIE'; the 14th visible column, since L, M, N and P are hidden).", False),
        ("T2: geur = D, stof = E, geluid = F, gevaar = I, grootste afstand = N, categorie = P. T2 contains no formulas; all values are typed.", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=1):
        readme.cell(row=i, column=1, value=txt).font = Font(name=FONT, bold=bold)
    readme.column_dimensions["A"].width = 160

    wb.save(OUT_XLSX)

    hashes_after = {k: sha256(src["path"]) for k, src in SOURCES.items()}
    assert hashes_before == hashes_after, "source files changed!"

    # Console summary for the documentation.
    for a in analyses:
        st = a["stats"]
        print(f"{a['label']}: n={st['n']} rho={st['rho']:.3f} [{a['ci'][0]:.3f},{a['ci'][1]:.3f}] "
              f"p={st['p_rho']:.2g} tau_b={st['tau']:.3f} p={st['p_tau']:.2g} r={st['r']:.3f} perm={a['perm']}")
    for key in SOURCES:
        inc = [r for r in extracts[key] if r["_d"]["included"]]
        reasons = {}
        for r in extracts[key]:
            reasons[r["_d"]["reason"]] = reasons.get(r["_d"]["reason"], 0) + 1
        print(key, "extracted", len(extracts[key]), "included", len(inc), "reasons", reasons,
              "O mismatches", [r["src_row"] for r in inc if not r["_d"]["O_ok"]],
              "R mismatches", [r["src_row"] for r in inc if not r["_d"]["R_ok"]],
              "R hardcoded", len([r for r in inc if not r["R_is_formula"]]),
              "decisive", sum(r["_d"]["decisive"] for r in inc),
              "geur 0", sum(r["_d"]["odour"] == 0 for r in inc),
              "in other", sum(r["in_both"] for r in datasets[key]))
    for r in t2_mismatch:
        print("T2 mismatch", r["src_row"], r["descr"], r["odour_raw"], r["dust_raw"], r["noise_raw"],
              r["danger_raw"], "max src", r["O_cached"], "recalc", r["_d"]["max_rec"],
              "cat src", r["R_cached"], "recalc", r["_d"]["cat_rec"])
    print("union n", len(union))


if __name__ == "__main__":
    main()
